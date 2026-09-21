"""Normalize user-supplied Bloomberg tables or screenshot transcriptions.

The script deliberately does not OCR images. Codex reads an attached screenshot,
transcribes only visible values to JSON, and sends that JSON through this helper.
CSV and JSON exports can be passed directly.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


ALIASES = {
    "symbol": ("symbol", "ticker", "security", "id_bb_global", "id_bb_sec_num_des"),
    "name": ("name", "security_des", "long_comp_name", "description"),
    "market": ("market", "country_iso", "country", "exchange"),
    "currency": ("currency", "crncy"),
    "last": ("last", "px_last", "last_price"),
    "prev_close": ("prev_close", "prev_close_value_realtime", "px_prev_close"),
    "pct_change": ("pct_change", "chg_pct_1d", "percent_change", "change_percent"),
    "open": ("open", "px_open"),
    "high": ("high", "px_high"),
    "low": ("low", "px_low"),
    "volume": ("volume", "px_volume", "volume_total"),
    "turnover": ("turnover", "turnover_total"),
    "sector": ("sector", "gics_sector_name", "bics_level_1_sector_name"),
    "benchmark": ("benchmark", "benchmark_index"),
    "timestamp": ("timestamp", "last_update_dt", "price_timestamp", "data_as_of"),
}

NUMERIC_FIELDS = {"last", "prev_close", "pct_change", "open", "high", "low", "volume", "turnover"}


def _key(value: Any) -> str:
    return str(value).strip().lower().replace(" ", "_").replace("-", "_")


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value: Any) -> float | int | None:
    if _blank(value):
        return None
    if isinstance(value, str):
        value = value.replace(",", "").replace("%", "").strip()
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(result):
        return None
    return int(result) if result.is_integer() else result


def normalize_rows(rows: list[dict[str, Any]], *, default_timestamp: str | None = None) -> tuple[list[dict[str, Any]], list[str]]:
    normalized = []
    warnings: list[str] = []
    for position, original in enumerate(rows, start=1):
        lookup = {_key(key): value for key, value in original.items()}
        row: dict[str, Any] = {}
        for field, aliases in ALIASES.items():
            value = next((lookup[alias] for alias in aliases if alias in lookup and not _blank(lookup[alias])), None)
            row[field] = _number(value) if field in NUMERIC_FIELDS else (str(value).strip() if not _blank(value) else None)
        if row["timestamp"] is None:
            row["timestamp"] = default_timestamp
        if row["pct_change"] is None and row["last"] is not None and row["prev_close"] not in (None, 0):
            row["pct_change"] = round((row["last"] / row["prev_close"] - 1) * 100, 6)
        identity = row["symbol"] or row["name"]
        if identity is None:
            warnings.append(f"row {position}: no symbol or name; row omitted")
            continue
        if row["last"] is None and row["pct_change"] is None:
            warnings.append(f"{identity}: neither last nor pct_change was visible")
        normalized.append({key: value for key, value in row.items() if value is not None})
    return normalized, warnings


def build_pack(
    rows: list[dict[str, Any]],
    *,
    source_kind: str,
    source_artifact: str | None = None,
    data_as_of: str | None = None,
    timezone_name: str | None = None,
    market: str | None = None,
    ingested_at: datetime | None = None,
) -> dict[str, Any]:
    normalized, warnings = normalize_rows(rows, default_timestamp=data_as_of)
    timestamps = sorted({str(row["timestamp"]) for row in normalized if row.get("timestamp")})
    return {
        "schemaVersion": 1,
        "provider": "Bloomberg user upload",
        "source_kind": source_kind,
        "source_artifact": source_artifact,
        "ingested_at": (ingested_at or datetime.now(timezone.utc)).isoformat(),
        "data_as_of": data_as_of,
        "timezone": timezone_name,
        "market": market,
        "movement_basis": "latest_vs_previous_close",
        "capture_window": {
            "start": timestamps[0] if timestamps else data_as_of,
            "end": timestamps[-1] if timestamps else data_as_of,
        },
        "quotes": normalized,
        "quality": {
            "row_count": len(normalized),
            "timestamp_verified": bool(timestamps or data_as_of),
            "warnings": warnings,
        },
    }


def load_rows(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle)), {}
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX input requires openpyxl from requirements.lock.txt.") from exc
        workbook = load_workbook(path, data_only=True, read_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        header_index = next((index for index, row in enumerate(values) if any(not _blank(value) for value in row)), None)
        if header_index is None:
            return [], {}
        headers = [str(value).strip() if not _blank(value) else f"column_{position}" for position, value in enumerate(values[header_index], start=1)]
        rows = [dict(zip(headers, row)) for row in values[header_index + 1:] if any(not _blank(value) for value in row)]
        return rows, {"worksheet": sheet.title}
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list):
        return data, {}
    if isinstance(data, dict):
        rows = data.get("quotes", data.get("rows"))
        if isinstance(rows, list):
            return rows, data
    raise ValueError("Input must be XLSX/CSV, a JSON row array, or an object containing quotes/rows.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", help="Bloomberg XLSX/CSV/JSON export, or JSON transcribed from a screenshot")
    parser.add_argument("--source-kind", choices=("bloomberg_export", "bloomberg_screenshot"), required=True)
    parser.add_argument("--data-as-of", help="Visible/exported market timestamp with timezone offset")
    parser.add_argument("--timezone")
    parser.add_argument("--market")
    parser.add_argument("--output")
    args = parser.parse_args()
    rows, metadata = load_rows(Path(args.input))
    pack = build_pack(
        rows,
        source_kind=args.source_kind,
        source_artifact=Path(args.input).name,
        data_as_of=args.data_as_of or metadata.get("data_as_of"),
        timezone_name=args.timezone or metadata.get("timezone"),
        market=args.market or metadata.get("market"),
    )
    rendered = json.dumps(pack, indent=2, ensure_ascii=False, allow_nan=False) + "\n"
    if args.output:
        Path(args.output).write_text(rendered, encoding="utf-8")
    else:
        print(rendered, end="")


if __name__ == "__main__":
    main()
